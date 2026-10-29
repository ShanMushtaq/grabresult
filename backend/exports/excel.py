"""XLSX export with openpyxl: multi-sheet workbook, filters, formatting."""
import json
import time
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADERS = [
    "Business Name", "Description", "Industry", "Category", "Website",
    "Country", "City", "Address", "Phone", "Email", "Socials",
    "Source", "Source URL", "Discovered", "Status", "Confidence",
    "Lead Score", "Score Reasons", "Notes",
]

HEADER_FILL = PatternFill("solid", fgColor="1F6FEB")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _row(lead: dict) -> list:
    try:
        socials = json.loads(lead.get("socials_json") or "{}")
    except Exception:
        socials = {}
    try:
        reasons = json.loads(lead.get("score_reasons_json") or "[]")
    except Exception:
        reasons = []
    socials_s = "; ".join(f"{k}: {v}" for k, v in socials.items())
    discovered = time.strftime("%Y-%m-%d %H:%M", time.localtime(lead.get("discovered_at") or 0))
    return [
        lead.get("business_name", ""), lead.get("description", ""),
        lead.get("industry", ""), lead.get("category", ""),
        lead.get("website", ""), lead.get("country", ""), lead.get("city", ""),
        lead.get("address", ""), lead.get("phone", ""), lead.get("email", ""),
        socials_s, lead.get("source_name", ""), lead.get("source_url", ""),
        discovered, lead.get("verification_status", ""),
        lead.get("confidence", ""), lead.get("lead_score", ""),
        "; ".join(reasons), lead.get("notes", ""),
    ]


def _style_sheet(ws, nrows):
    for col in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(nrows + 1, 2)}"
    widths = [28, 40, 18, 18, 32, 16, 18, 34, 18, 28, 40, 22, 40, 17, 14, 11, 11, 50, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _add_sheet(wb, title, leads):
    ws = wb.create_sheet(title[:31])
    ws.append(HEADERS)
    for lead in leads:
        ws.append(_row(lead))
    _style_sheet(ws, len(leads))
    return ws


def build_workbook(leads: list, summary: dict, source_log: list, path: str) -> dict:
    wb = Workbook()
    wb.remove(wb.active)

    _add_sheet(wb, "All Results", leads)
    _add_sheet(wb, "Qualified Leads",
               [l for l in leads if l.get("verification_status") == "qualified"])
    _add_sheet(wb, "Needs Review",
               [l for l in leads if l.get("verification_status") == "needs_review"])

    by_country: dict[str, list] = {}
    by_city: dict[str, list] = {}
    for l in leads:
        by_country.setdefault(l.get("country") or "Unknown", []).append(l)
        by_city.setdefault(l.get("city") or "Unknown", []).append(l)
    for country in sorted(by_country)[:20]:
        _add_sheet(wb, f"Country - {country}"[:31], by_country[country])
    for city in sorted(by_city)[:20]:
        _add_sheet(wb, f"City - {city}"[:31], by_city[city])

    ws = wb.create_sheet("Sources and Evidence")
    ws.append(["Source", "Status", "Message"])
    for s in source_log:
        ws.append([s.get("source", ""), s.get("status", ""), s.get("message", "")])
    _style_sheet(ws, len(source_log))

    ws = wb.create_sheet("Research Summary")
    ws.append(["Metric", "Value"])
    for k, v in (summary or {}).items():
        if isinstance(v, dict):
            for sk, sv in v.items():
                ws.append([f"{k}: {sk}", sv])
        else:
            ws.append([k, v])
    _style_sheet(ws, 30)

    wb.save(path)
    return {"sheets": wb.sheetnames, "rows": len(leads)}
